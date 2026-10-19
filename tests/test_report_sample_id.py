"""S16 regression: the statistical report keeps Sample_ID.

The per-sample category sheets only include the Sample ID column when
ExcelReportGenerator receives ``sample_id_col``; app.py did not pass it.
Synthetic data only.
"""
import ast
import re
import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from modules.report_generation import ExcelReportGenerator

ROOT = Path(__file__).resolve().parents[1]
BAS = ["CA", "CDCA", "DCA", "LCA", "GCA", "TCA", "GCDCA", "TCDCA"]


def _synthetic():
    rng = np.random.default_rng(7)
    df = pd.DataFrame({"Sample ID": [f"S{i:02d}" for i in range(12)],
                       "Group": np.repeat(["A", "B"], 6)})
    for b in BAS:
        df[b] = np.round(rng.lognormal(1.0, 0.5, 12), 3)
    return df


def _ids_per_sheet(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return {s: sum(1 for row in wb[s].iter_rows(values_only=True) for x in row
                       if isinstance(x, str) and re.fullmatch(r"S\d\d", x))
                for s in wb.sheetnames if s != "Overview"}
    finally:
        wb.close()


def _report(tmp_path, sample_id_col):
    gen = ExcelReportGenerator(data=_synthetic(), group_col="Group",
                               sample_id_col=sample_id_col, bile_acid_cols=BAS)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        gen.run_section("categories")
        out = tmp_path / "report.xlsx"
        gen.save_excel_report(out)
    return _ids_per_sheet(out)


def test_category_sheets_carry_sample_id(tmp_path):
    hits = _report(tmp_path, "Sample ID")
    assert hits, "no category sheets written"
    assert all(n == 12 for n in hits.values()), hits


def test_sample_id_equal_to_group_col_does_not_duplicate(tmp_path):
    gen = ExcelReportGenerator(data=_synthetic(), group_col="Group",
                               sample_id_col="Group", bile_acid_cols=BAS)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        gen.run_section("categories")
    sheet = next(iter(gen.analysis_sheets.values()))
    assert list(sheet.raw_data.columns).count("Group") == 1


def test_app_passes_sample_id_col_to_report_generator():
    tree = ast.parse((ROOT / "app.py").read_text(encoding="utf-8"))
    calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call)
             and getattr(n.func, "id", getattr(n.func, "attr", None)) == "ExcelReportGenerator"]
    assert calls, "ExcelReportGenerator call not found in app.py"
    for call in calls:
        assert "sample_id_col" in {k.arg for k in call.keywords}
