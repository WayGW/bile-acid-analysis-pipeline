"""User-defined group order (feature/group-order). Synthetic data only."""
import numpy as np
import pandas as pd
import pytest

from modules.group_order import (
    apply_order, combined_order, default_levels, level_key, order_table, positions_to_order,
)


# ---------------------------------------------------------------- helpers
def test_apply_order_full_and_none():
    obs = ["Control", "DDC", "DDC + RiVM"]
    assert apply_order(obs, None) == obs
    assert apply_order(obs, []) == obs
    assert apply_order(obs, ["DDC + RiVM", "Control", "DDC"]) == ["DDC + RiVM", "Control", "DDC"]


def test_apply_order_partial_and_unseen_levels():
    obs = ["A", "B", "C", "D"]
    # named levels first in user order, the rest keep their default order
    assert apply_order(obs, ["C", "A"]) == ["C", "A", "B", "D"]
    # preferred levels absent from this panel are skipped
    assert apply_order(["B", "D"], ["D", "X", "B"]) == ["D", "B"]


def test_apply_order_matches_on_stripped_labels_and_returns_observed_objects():
    obs = ["Control", "DDC + Ang 1-7 ", "DDC"]
    out = apply_order(obs, ["DDC + Ang 1-7", "DDC", "Control"])
    assert out == ["DDC + Ang 1-7 ", "DDC", "Control"]      # original value kept
    assert out[0] is obs[1]


def test_apply_order_numeric_levels():
    assert apply_order([1, 2, 3], ["3", "1"]) == [3, 1, 2]


def test_default_levels_modes():
    s = pd.Series(["b", "a", np.nan, "b", "c", "nan"])
    assert default_levels(s, "appearance") == ["b", "a", "c"]
    assert default_levels(s, "sorted") == ["a", "b", "c"]
    with pytest.raises(ValueError):
        default_levels(s, "other")


def test_level_key():
    assert level_key(" x ") == "x"
    assert level_key(np.nan) is None
    assert level_key(None) is None
    assert level_key(2) == "2"


def test_order_table_and_round_trip():
    t = order_table(["Control", "DDC"])
    assert t.to_dict("list") == {"Group": ["Control", "DDC"], "Position": [1, 2]}
    t["Position"] = [2, 1]
    assert positions_to_order(t) == (["DDC", "Control"], None)


def test_positions_allow_gaps_reject_duplicates_and_fractions():
    t = pd.DataFrame({"Group": ["A", "B", "C"], "Position": [7, 1, 3]})
    assert positions_to_order(t) == (["B", "C", "A"], None)
    order, err = positions_to_order(pd.DataFrame({"Group": ["A", "B"], "Position": [1, 1]}))
    assert order is None and "repeated: 1" in err
    order, err = positions_to_order(pd.DataFrame({"Group": ["A", "B"], "Position": [1, 1.5]}))
    assert order is None and "whole numbers" in err
    order, err = positions_to_order(pd.DataFrame({"Group": ["A", "B"], "Position": [1, None]}))
    assert order is None and "B" in err


def test_combined_order_matches_app_labels():
    # app.py builds labels as A.astype(str) + ' - ' + B.astype(str) [+ ' - ' + C]
    assert combined_order([["WT", "KO"], ["Chow", "WD"]]) == [
        "WT - Chow", "WT - WD", "KO - Chow", "KO - WD"]
    assert combined_order([["x"], ["y"], ["1", "2"]]) == ["x - y - 1", "x - y - 2"]
    assert combined_order([]) == []


# ---------------------------------------------------------------- figures
import ast
import contextlib
import io
import warnings
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from config.bile_acid_species import ANALYSIS_GROUP_DISPLAY_NAMES
from modules.report_generation import ExcelReportGenerator
from modules.statistical_tests import StatisticalAnalyzer, significant_pairs
from modules.visualization import BileAcidVisualizer

ROOT = Path(__file__).resolve().parents[1]
GROUPS = ["Ctrl", "Low", "High"]
FACTOR_ORDERS = {"Factor_Geno": ["WT", "KO"], "Factor_Diet": ["WD", "Chow"], "Factor_Tx": ["Veh", "Drug"]}


@pytest.fixture(scope="module")
def oneway():
    rng = np.random.default_rng(3)
    df = pd.DataFrame({
        "Group": np.repeat(GROUPS, 8),
        "CA": np.concatenate([rng.normal(m, 0.3, 8) for m in (1, 3, 6)]),
        "DCA": np.concatenate([rng.normal(m, 0.3, 8) for m in (2, 2.2, 5)]),
    })
    res = {c: StatisticalAnalyzer().analyze(df, c, "Group") for c in ("CA", "DCA")}
    return df, res


@pytest.fixture(scope="module")
def factorial():
    rng = np.random.default_rng(5)
    df = pd.DataFrame([(a, b, c) for a in ("WT", "KO") for b in ("Chow", "WD") for c in ("Veh", "Drug")
                       for _ in range(4)], columns=["Factor_Geno", "Factor_Diet", "Factor_Tx"])
    df["CA"] = rng.lognormal(1.0, 0.4, len(df))
    df["DCA"] = rng.lognormal(1.5, 0.4, len(df))
    df["Sample ID"] = [f"s{i:02d}" for i in range(len(df))]
    kw = dict(data=df, group_col="Factor_Geno", bile_acid_cols=["CA", "DCA"], n_samples=len(df))
    tw = ExcelReportGenerator(factors={"Geno": "Factor_Geno", "Diet": "Factor_Diet"}, n_factors=2, **kw)
    th = ExcelReportGenerator(factors={"Geno": "Factor_Geno", "Diet": "Factor_Diet", "Tx": "Factor_Tx"},
                              n_factors=3, **kw)
    with contextlib.redirect_stdout(io.StringIO()), warnings.catch_warnings():
        warnings.simplefilter("ignore")
        r_tw = tw.run_section("individual_ba").twoway_individual_ba
        r_th = th.run_section("individual_ba").threeway_individual_ba
    return df, r_tw, r_th


def _xt(ax):
    return [t.get_text() for t in ax.get_xticklabels()]


def test_oneway_ticks_follow_order_and_default_is_unchanged(oneway):
    df, res = oneway
    fig = BileAcidVisualizer().plot_multi_panel_groups_with_stats(df, ["CA", "DCA"], "Group", res, ncols=2)
    assert _xt(fig.axes[0]) == GROUPS          # first appearance, as before
    plt.close(fig)
    order = ["High", "Ctrl", "Low"]
    fig = BileAcidVisualizer(level_orders={"Group": order}).plot_multi_panel_groups_with_stats(
        df, ["CA", "DCA"], "Group", res, ncols=2)
    assert _xt(fig.axes[0]) == order
    plt.close(fig)


def test_brackets_sit_on_reordered_groups(oneway):
    df, res = oneway
    order = ["High", "Ctrl", "Low"]
    pos = {g: i for i, g in enumerate(order)}
    fig = BileAcidVisualizer(level_orders={"Group": order}).plot_multi_panel_groups_with_stats(
        df, ["CA"], "Group", res, ncols=1, plot_type="box", max_sig_bars=3)
    ax = fig.axes[0]
    drawn = sorted(tuple(sorted({int(x) for x in ln.get_xdata()})) for ln in ax.lines if len(ln.get_xdata()) == 4)
    expected = sorted(tuple(sorted((pos[a], pos[b]))) for a, b, _ in significant_pairs(res["CA"], max_pairs=3))
    assert expected and drawn == expected
    plt.close(fig)


def test_twoway_axis_and_legend_follow_factor_orders(factorial):
    df, r_tw, _ = factorial
    fig = BileAcidVisualizer(level_orders=FACTOR_ORDERS).plot_twoway_multi_panel(
        df, ["CA"], "Factor_Geno", "Factor_Diet", r_tw, ncols=1, plot_type="box")
    ax = fig.axes[0]
    assert _xt(ax) == ["WT", "KO"]
    assert [t.get_text() for t in ax.get_legend().get_texts()] == ["WD", "Chow"]
    plt.close(fig)
    fig = BileAcidVisualizer().plot_twoway_multi_panel(
        df, ["CA"], "Factor_Geno", "Factor_Diet", r_tw, ncols=1, plot_type="box")
    assert _xt(fig.axes[0]) == ["KO", "WT"]    # alphabetical default, as before
    plt.close(fig)


def test_threeway_facets_axis_and_legend_follow_factor_orders(factorial):
    df, _, r_th = factorial
    fig = BileAcidVisualizer(level_orders=FACTOR_ORDERS).plot_threeway_multi_panel(
        df, ["CA"], "Factor_Geno", "Factor_Diet", "Factor_Tx", r_th, ncols=2)
    titles = [a.get_title() for a in fig.axes if a.get_title()]
    assert titles[:2] == ["Factor_Tx = Veh", "Factor_Tx = Drug"]
    assert _xt(fig.axes[0]) == ["WT", "KO"]
    assert [t.get_text() for t in fig.legends[0].get_texts()] == ["WD", "Chow"]
    plt.close(fig)


def test_stacked_composition_follows_order(factorial):
    df, _, _ = factorial
    out = BileAcidVisualizer(level_orders={"Factor_Geno": ["WT", "KO"]}).plot_composition_stacked(
        df, "Factor_Geno", ["CA", "DCA"])
    fig = out[0] if isinstance(out, tuple) else out
    assert _xt(fig.axes[0]) == ["WT", "KO"]
    plt.close(fig)


def test_heatmap_rows_and_legend_follow_factor_orders(factorial):
    df, _, _ = factorial
    fc = {"Geno": "Factor_Geno", "Diet": "Factor_Diet", "Tx": "Factor_Tx"}
    fig = BileAcidVisualizer(level_orders=FACTOR_ORDERS).plot_faceted_heatmap(
        df, pd.DataFrame({"total_primary": df.CA + df.DCA}), "Factor_Geno",
        {"total_primary": ["CA", "DCA"]}, ANALYSIS_GROUP_DISPLAY_NAMES,
        sample_id_col="Sample ID", factor_cols=fc)
    rows = [t.get_text() for t in fig.axes[0].get_yticklabels()]
    first = df[(df.Factor_Geno == "WT") & (df.Factor_Diet == "WD") & (df.Factor_Tx == "Veh")]["Sample ID"].tolist()
    assert rows[:len(first)] == first            # WT / WD / Veh block first, sheet order inside
    legend = [t.get_text() for t in fig.legends[0].get_texts()]
    assert legend == ["Geno: WT", "Geno: KO", "Diet: WD", "Diet: Chow", "Tx: Veh", "Tx: Drug"]
    plt.close(fig)


def test_report_statistics_unaffected_by_display_orders(factorial, tmp_path):
    df, _, _ = factorial
    gen = ExcelReportGenerator(data=df, group_col="Factor_Geno", bile_acid_cols=["CA", "DCA"],
                               factors={"Geno": "Factor_Geno", "Diet": "Factor_Diet"}, n_factors=2,
                               n_samples=len(df))
    with contextlib.redirect_stdout(io.StringIO()), warnings.catch_warnings():
        warnings.simplefilter("ignore")
        gen.run_all_statistics()
        gen.save_excel_report(tmp_path / "a.xlsx")
        gen.display_orders = {"Factor_Diet": ["WD", "Chow"]}
        gen.save_excel_report(tmp_path / "b.xlsx")
    import openpyxl

    def sheets(p):
        wb = openpyxl.load_workbook(p, read_only=True)
        try:
            return {s: [tuple(r) for r in wb[s].iter_rows(values_only=True)] for s in wb.sheetnames}
        finally:
            wb.close()
    a, b = sheets(tmp_path / "a.xlsx"), sheets(tmp_path / "b.xlsx")
    assert list(a) == list(b)
    assert all(a[s] == b[s] for s in a if s != "Overview")
    flat = [str(c) for r in b["Overview"] for c in r if c is not None]
    assert "FIGURE GROUP ORDER (display only; statistics unaffected)" in flat
    assert "WD | Chow" in flat
    assert not any("FIGURE GROUP ORDER" in str(c) for r in a["Overview"] for c in r if c is not None)


# ---------------------------------------------------------------- app wiring
def test_app_builds_every_visualizer_through_make_visualizer():
    tree = ast.parse((ROOT / "app.py").read_text(encoding="utf-8"))
    owners = []
    for fn in [n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef)]:
        for call in [n for n in ast.walk(fn) if isinstance(n, ast.Call)]:
            if getattr(call.func, "id", getattr(call.func, "attr", None)) == "BileAcidVisualizer":
                owners.append(fn.name)
                assert "level_orders" in {k.arg for k in call.keywords}
    assert owners == ["make_visualizer"]


_APP_SCRIPT = '''
import sys, importlib.util
import pandas as pd
from types import SimpleNamespace
sys.path.insert(0, r"{root}")
spec = importlib.util.spec_from_file_location("bile_app", r"{app}")
app = importlib.util.module_from_spec(spec); spec.loader.exec_module(app)
import streamlit as st
app.init_session_state()
if "sentinel_set" not in st.session_state:
    st.session_state.sentinel_set = True
    st.session_state.stats_computed = True
    st.session_state.report_generator = "cached-generator"
    st.session_state.heatmap_fig = b"old-png"
sd = pd.DataFrame({{"Type": ["Ctrl", "Ctrl", "DDC", "DDC", "DDC + Ang 1-7 ", "DDC + Ang 1-7 "], "CA": range(6)}})
processed = SimpleNamespace(sample_data=sd, structure=SimpleNamespace(group_col="Type", n_factors=0, factors={{}}))
app.render_group_order_controls(processed)
'''


def _apptest():
    testing = pytest.importorskip("streamlit.testing.v1")
    at = testing.AppTest.from_string(_APP_SCRIPT.format(root=ROOT, app=ROOT / "app.py"), default_timeout=60)
    at.run()
    assert not at.exception
    return at


def _click(at, label):
    [b for b in at.button if b.label == label][0].click()
    at.run()


def test_ui_apply_reorders_and_keeps_statistics_cached():
    at = _apptest()
    at.session_state["group_order_editor_Type"] = {
        "edited_rows": {0: {"Position": 3}, 2: {"Position": 1}}, "added_rows": [], "deleted_rows": []}
    _click(at, "Apply order")
    assert not at.exception
    assert at.session_state["level_orders"] == {"Type": ["DDC + Ang 1-7 ", "DDC", "Ctrl"]}
    assert at.session_state["stats_computed"] is True
    assert at.session_state["report_generator"] == "cached-generator"
    assert "heatmap_fig" not in at.session_state      # rendered figure cache is dropped
    _click(at, "Reset to default")
    assert at.session_state["level_orders"] == {}


def test_ui_rejects_duplicate_positions():
    at = _apptest()
    at.session_state["group_order_editor_Type"] = {
        "edited_rows": {0: {"Position": 2}}, "added_rows": [], "deleted_rows": []}
    _click(at, "Apply order")
    assert at.session_state["level_orders"] == {}
    assert any("repeated: 2" in w.value for w in at.warning)
    assert at.session_state["heatmap_fig"] == b"old-png"   # nothing invalidated
