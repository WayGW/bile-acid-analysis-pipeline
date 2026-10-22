"""Regression tests for the v4.7.0 statistical-testing changes.

Covers aggregate-aware LOD exclusion, default-on cross-analyte FDR gating,
category-result reuse, and the two-way ART main-effect denominator fix.
Synthetic data only.
"""
import contextlib
import io

import numpy as np
import openpyxl
import pandas as pd
import pytest

from modules.report_generation import ExcelReportGenerator
from modules.statistical_tests import (
    StatisticalAnalyzer,
    TestType as StatTestType,
    adjust_pvalues_bh,
    apply_fdr_to_oneway_results,
    apply_fdr_to_twoway_results,
)


BAS = ["CA", "CDCA", "DCA"]


def _lod_data():
    """Group A is LOD-bound for primary BAs; group B is fully observed."""
    df = pd.DataFrame({
        "Sample ID": [f"S{i}" for i in range(8)],
        "Group": ["A"] * 4 + ["B"] * 4,
        "CA": [5.0, 5.0, 5.0, 5.0, 10.0, 12.0, 14.0, 16.0],
        "CDCA": [5.0, 5.0, 5.0, 5.0, 8.0, 9.0, 11.0, 12.0],
        "DCA": [1.0, 1.2, 1.4, 1.6, 20.0, 22.0, 24.0, 26.0],
    })
    return df


def _generator(df, **kwargs):
    lod_rows = {"CA": list(range(4)), "CDCA": list(range(4))}
    lod_values = {"CA": 5.0, "CDCA": 5.0, "DCA": 1.0}
    totals = pd.DataFrame({
        "total_primary": df[["CA", "CDCA"]].sum(axis=1),
        "total_secondary": df["DCA"],
        "total_unconjugated": df[BAS].sum(axis=1),
        "total_all": df[BAS].sum(axis=1),
    })
    params = dict(
        data=df,
        group_col="Group",
        sample_id_col="Sample ID",
        bile_acid_cols=BAS,
        totals=totals,
        n_samples=len(df),
        analyte_lod_rows=lod_rows,
        analyte_lod_values=lod_values,
        aggregate_lod_threshold=5.0,
        fdr_enabled=False,
    )
    params.update(kwargs)
    return ExcelReportGenerator(**params)


def _run(gen, section):
    with contextlib.redirect_stdout(io.StringIO()):
        return gen.run_section(section)


def test_total_all_is_warning_only_but_other_aggregates_are_excluded():
    gen = _generator(_lod_data())
    results = _run(gen, "totals")

    # total_all is always tested even when a group is 100% LOD-bound.
    assert "total_all" in results.totals_results
    assert "total_all" not in results.lod_excluded
    assert "total_all" in results.lod_warnings
    assert results.lod_warnings["total_all"]["max_group_burden_pct"] >= 5.0

    # Primary total is fully LOD-bound in group A and is excluded.
    assert "total_primary" in results.lod_excluded
    assert "total_primary" not in results.totals_results
    detail = results.aggregate_lod_details["total_primary"]
    assert detail["excluded"] is True
    assert detail["group_fractions"]["A"] == pytest.approx(1.0)
    assert len(detail["sample_frame"]) == len(_lod_data())
    assert {"lod_bound_contribution", "observed_contribution", "lod_burden_pct"}.issubset(
        detail["sample_frame"].columns
    )


def test_aggregate_lod_uses_burden_not_constituent_exclusion_count():
    """All primary BAs are LOD-bound, so burden is 100% in group A."""
    gen = _generator(_lod_data())
    info = gen._aggregate_lod_burden_info("total_primary")

    assert info is not None
    assert info["max_group_fraction"] == pytest.approx(1.0)
    # Per-sample burden is retained, including the observed group.
    per_sample = info["sample_frame"]
    assert per_sample["lod_burden_pct"].max() == pytest.approx(100.0)
    assert per_sample["lod_burden_pct"].min() == pytest.approx(0.0)


def test_category_sheets_reuse_totals_result_and_inherit_lod_note():
    gen = _generator(_lod_data(), fdr_enabled=True)
    _run(gen, "totals")
    totals_result = gen.results.totals_results["total_all"]
    n_families_before = dict(gen.results.fdr_families)

    results = _run(gen, "categories")

    sheet = gen.analysis_sheets["Total_All_BAs"]
    assert sheet.statistical_result is not None
    assert sheet.statistical_result.main_test.pvalue == pytest.approx(
        totals_result.main_test.pvalue
    )
    assert sheet.statistical_result.main_test.qvalue == pytest.approx(
        totals_result.main_test.qvalue
    )
    assert "Retained with warning" in sheet.statistical_note
    assert results.category_results["Total_All_BAs"] is sheet.statistical_result

    excluded_sheet = gen.analysis_sheets["Total_Primary"]
    assert excluded_sheet.statistical_result is None
    assert "Not tested" in excluded_sheet.statistical_note

    # Category sheets must not create duplicate FDR families or tests.
    assert gen.results.fdr_families == n_families_before


def test_excel_report_includes_per_sample_aggregate_burden(tmp_path):
    gen = _generator(_lod_data())
    _run(gen, "totals")
    path = tmp_path / "report.xlsx"
    gen.save_excel_report(path)

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        assert "Aggregate_LOD_Burden" in wb.sheetnames
        rows = list(wb["Aggregate_LOD_Burden"].iter_rows(values_only=True))
    finally:
        wb.close()

    flat = [value for row in rows for value in row if value is not None]
    for sample_id in ("S0", "S7"):
        assert sample_id in flat
    assert "Retained (warning)" in flat
    assert "Excluded" in flat


def test_bh_helper_preserves_nans_and_controls_family_size():
    pvalues = [0.01, np.nan, 0.04, 0.03]
    adjusted = adjust_pvalues_bh(pvalues)

    assert np.isnan(adjusted[1])
    finite = adjusted[[0, 2, 3]]
    assert np.all(finite >= np.array([0.01, 0.04, 0.03]) - 1e-12)
    assert np.all(finite <= 1.0)


def _fdr_data():
    rng = np.random.default_rng(13)
    groups = np.repeat(["A", "B", "C"], 8)
    return pd.DataFrame({
        "Group": groups,
        "Strong1": np.concatenate([
            rng.normal(0.0, 0.15, 8),
            rng.normal(1.5, 0.15, 8),
            rng.normal(3.0, 0.15, 8),
        ]),
        "Strong2": np.concatenate([
            rng.normal(0.0, 0.20, 8),
            rng.normal(1.2, 0.20, 8),
            rng.normal(2.4, 0.20, 8),
        ]),
        "Null": rng.normal(1.0, 0.5, 24),
    })


def test_oneway_fdr_is_default_on_separates_scales_and_gates_posthoc():
    gen = ExcelReportGenerator(
        data=_fdr_data(),
        group_col="Group",
        bile_acid_cols=["Strong1", "Strong2", "Null"],
        n_samples=24,
    )
    results = _run(gen, "individual_ba")

    # FDR is on unless explicitly disabled.
    assert results.fdr_enabled is True
    strong = results.individual_ba_results["Strong1"]
    assert strong.main_test.qvalue is not None
    assert strong.main_test.qvalue >= strong.main_test.pvalue - 1e-12
    assert strong.main_test.fdr_family == "oneway:individual_ba:raw"
    assert strong.main_test.significant == (strong.main_test.qvalue < gen.alpha)

    log_strong = results.log_individual_ba_results["Strong1"]
    assert log_strong.main_test.fdr_family == "oneway:individual_ba:log10"

    assert set(results.fdr_families) >= {
        "oneway:individual_ba:raw",
        "oneway:individual_ba:log10",
    }

    for result in list(results.individual_ba_results.values()) + list(
        results.log_individual_ba_results.values()
    ):
        if not result.main_test.significant:
            assert result.posthoc_test is None
        else:
            assert result.posthoc_test is not None

    # The null analyte should not acquire a post-hoc table merely because it
    # had a raw p-value below alpha.
    null = results.individual_ba_results.get("Null")
    if null is not None and not null.main_test.significant:
        assert null.posthoc_test is None


def test_twoway_fdr_is_separate_by_effect():
    rng = np.random.default_rng(17)
    df = pd.DataFrame([
        (a, b) for a in ("WT", "KO") for b in ("Chow", "WD") for _ in range(6)
    ], columns=["A", "B"])
    for name, shift in (("X", 1.0), ("Y", 0.4), ("Z", 0.0)):
        values = rng.normal(0, 1, len(df))
        values[df["A"].eq("KO")] += shift
        values[df["B"].eq("WD")] += shift / 2
        df[name] = values

    analyzer = StatisticalAnalyzer()
    results = {}
    for name in ("X", "Y", "Z"):
        results[name] = analyzer.analyze_twoway(
            df, name, "A", "B", "A", "B", run_posthoc=False
        )

    sizes = apply_fdr_to_twoway_results(results, 0.05, "twoway:test")

    assert set(sizes) == {
        "twoway:test:A",
        "twoway:test:B",
        "twoway:test:interaction",
    }
    for effect, q_attr in (
        ("A", "factor_a_qvalue"),
        ("B", "factor_b_qvalue"),
        ("interaction", "interaction_qvalue"),
    ):
        assert sizes[f"twoway:test:{effect}"] == 3
        assert all(
            getattr(result.twoway_result, q_attr) is not None
            for result in results.values()
        )


def test_twoway_art_main_effects_use_full_factorial_error_df():
    rng = np.random.default_rng(19)
    df = pd.DataFrame([
        (a, b) for a in ("WT", "KO") for b in ("Chow", "WD") for _ in range(6)
    ], columns=["A", "B"])
    # Skewed/mixture data pushes the analyzer toward ART.
    df["Y"] = np.concatenate([
        rng.exponential(1.0, 6),
        rng.exponential(2.0, 6),
        rng.exponential(4.0, 6),
        rng.exponential(3.0, 6),
    ])

    analyzer = StatisticalAnalyzer()
    result = analyzer._run_art_anova(df, "Y", "A", "B", "A", "B")

    assert result.test_type == StatTestType.ART_ANOVA
    n = len(df)
    cells = df["A"].nunique() * df["B"].nunique()
    expected_residual_df = n - cells

    # The old bug used one-way error df (n - n_levels) for main effects.
    assert result.factor_a_df[1] == expected_residual_df
    assert result.factor_b_df[1] == expected_residual_df
    assert result.interaction_df[1] == expected_residual_df

    # Main-effect F statistics must come from the full factorial model, not
    # from a one-way F on the same aligned ranks.
    from scipy.stats import f_oneway, rankdata

    values = df["Y"].to_numpy(dtype=float)
    grand = values.mean()
    marginal_a = df.groupby("A")["Y"].mean().to_dict()
    marginal_b = df.groupby("B")["Y"].mean().to_dict()
    cell_means = df.groupby(["A", "B"])["Y"].mean().to_dict()
    aligned = np.array([
        y - (marginal_b[b] - grand)
        - (cell_means[(a, b)] - marginal_a[a] - marginal_b[b] + grand)
        for y, a, b in zip(values, df["A"], df["B"])
    ])
    ranks = rankdata(aligned)
    oneway_f = f_oneway(
        *[ranks[df["A"].to_numpy() == level] for level in sorted(df["A"].unique())]
    ).statistic

    assert np.isfinite(result.factor_a_stat)
    assert np.isfinite(result.factor_b_stat)
    assert not np.isclose(result.factor_a_stat, oneway_f)
